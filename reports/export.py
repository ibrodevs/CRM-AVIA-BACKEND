"""Выгрузка отчёта в XLSX и CSV."""

from __future__ import annotations

import csv
import io
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_COLUMNS = (
    ("group", "Группа"),
    ("currency", "Валюта"),
    ("services", "Услуг"),
    ("revenue", "Выручка"),
    ("cost", "Себестоимость"),
    ("profit", "Прибыль"),
    ("average_check", "Средний чек"),
    ("margin_percent", "Маржа, %"),
)

_GROUP_TITLES = {
    "period": "Период",
    "operator": "Оператор",
    "supplier": "Поставщик",
    "kind": "Вид услуг",
    "client": "Клиент",
}


def _header(report: dict) -> list[str]:
    titles = [title for _, title in _COLUMNS]
    titles[0] = _GROUP_TITLES.get(report["group_by"], "Группа")
    return titles


def to_xlsx(report: dict) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"

    sheet["A1"] = f"Отчёт: {_GROUP_TITLES.get(report['group_by'], report['group_by'])}"
    sheet["A1"].font = Font(bold=True, size=13)
    filters = report["filters"]
    sheet["A2"] = f"Период: {filters.get('from') or 'с начала'} — {filters.get('to') or 'по сегодня'}"

    header_row = 4
    for index, title in enumerate(_header(report), start=1):
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = Font(bold=True)

    row_index = header_row + 1
    for row in report["rows"]:
        for column_index, (key, _) in enumerate(_COLUMNS, start=1):
            value = row[key]
            if key in ("revenue", "cost", "profit", "average_check", "margin_percent"):
                value = float(value)
            sheet.cell(row=row_index, column=column_index, value=value)
        row_index += 1

    row_index += 1
    for total in report["totals"]:
        sheet.cell(row=row_index, column=1, value="Итого").font = Font(bold=True)
        sheet.cell(row=row_index, column=2, value=total["currency"]).font = Font(bold=True)
        sheet.cell(row=row_index, column=3, value=total["services"]).font = Font(bold=True)
        sheet.cell(row=row_index, column=4, value=float(total["revenue"])).font = Font(bold=True)
        sheet.cell(row=row_index, column=5, value=float(total["cost"])).font = Font(bold=True)
        sheet.cell(row=row_index, column=6, value=float(total["profit"])).font = Font(bold=True)
        row_index += 1

    for index in range(1, len(_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 20

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(report: dict) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")
    writer.writerow(_header(report))
    for row in report["rows"]:
        writer.writerow([row[key] for key, _ in _COLUMNS])
    writer.writerow([])
    for total in report["totals"]:
        writer.writerow(
            ["Итого", total["currency"], total["services"], total["revenue"], total["cost"], total["profit"]]
        )
    # BOM — чтобы Excel открыл кириллицу без ручного выбора кодировки.
    return "﻿" + buffer.getvalue()
